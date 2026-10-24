import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from docx import Document

EMAIL = re.compile(r'(?<![\w.+-])[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}')
URL = re.compile(r'(?i)\b(?:https?://|www\.)[^\s<>]+')


def load_dictionary(path):
    with open(path, encoding='utf-8-sig', newline='') as stream:
        rows = csv.DictReader(stream)
        if not {'original', 'replacement'}.issubset(rows.fieldnames or []):
            raise ValueError('Dictionary needs original,replacement columns')
        mapping = {}
        for row in rows:
            original, replacement = row['original'].strip(), row['replacement'].strip()
            if original and replacement:
                mapping[original] = replacement
    return mapping


def process(input_file, dictionary_file, output_dir):
    source = Path(input_file)
    if source.suffix.lower() != '.docx':
        raise ValueError('Version 0.1 supports .docx files only')
    target_dir = Path(output_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    mapping = load_dictionary(dictionary_file)
    patterns = [(re.compile(r'(?<!\w)' + re.escape(k) + r'(?!\w)', re.I), k, v) for k, v in sorted(mapping.items(), key=lambda item: -len(item[0]))]
    doc = Document(source)
    events = []

    def replace_paragraph(paragraph, location):
        original = paragraph.text
        changed = original
        for pattern, key, value in patterns:
            changed, count = pattern.subn(lambda _: value, changed)
            if count:
                events.append({'location': location, 'original': key, 'replacement': value, 'count': count})
        for pattern, marker in ((EMAIL, '[Email removed]'), (URL, '[URL removed]')):
            matches = pattern.findall(changed)
            for match in matches:
                events.append({'location': location, 'original': match, 'replacement': marker, 'count': 1})
            changed = pattern.sub(lambda _: marker, changed)
        if changed != original:
            # Replacing paragraph text may remove inline formatting. This is a known v0.1 limitation.
            paragraph.text = changed

    def walk(container, prefix):
        for i, paragraph in enumerate(container.paragraphs, 1):
            replace_paragraph(paragraph, f'{prefix}/paragraph-{i}')
        for ti, table in enumerate(container.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                for ci, cell in enumerate(row.cells, 1):
                    walk(cell, f'{prefix}/table-{ti}/row-{ri}/cell-{ci}')

    walk(doc, 'body')
    for i, section in enumerate(doc.sections, 1):
        walk(section.header, f'header-{i}')
        walk(section.footer, f'footer-{i}')
    output_file = target_dir / f'{source.stem}_deidentified.docx'
    if output_file.resolve() == source.resolve():
        raise ValueError('Output must differ from original')
    doc.save(output_file)
    log_file = target_dir / f'{source.stem}_audit.csv'
    with log_file.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['location', 'original', 'replacement', 'count'])
        writer.writeheader()
        writer.writerows(events)
    manifest = {
        'software_version': '0.1.0',
        'processed_utc': datetime.now(timezone.utc).isoformat(),
        'input_name': source.name,
        'input_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'dictionary_sha256': hashlib.sha256(Path(dictionary_file).read_bytes()).hexdigest(),
        'output_name': output_file.name,
        'replacement_events': sum(event['count'] for event in events),
        'limitations': ['Review all output manually', 'Inline formatting may be lost in modified paragraphs', 'Embedded hyperlinks, images, comments and metadata are not fully sanitized']
    }
    (target_dir / f'{source.stem}_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
    return output_file, log_file, manifest
