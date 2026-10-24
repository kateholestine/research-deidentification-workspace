from docx import Document
from rdw.engine import process


def test_replacements(tmp_path):
    source = tmp_path / 'input.docx'
    doc = Document()
    doc.add_paragraph('NIC contacted NWCCU at person@example.org')
    doc.save(source)
    dictionary = tmp_path / 'dictionary.csv'
    dictionary.write_text('original,replacement\nNIC,Community College A\nNWCCU,Regional Accrediting Agency A\n', encoding='utf-8')
    output, log, manifest = process(source, dictionary, tmp_path / 'out')
    text = Document(output).paragraphs[0].text
    assert 'Community College A' in text
    assert '[Email removed]' in text
    assert manifest['replacement_events'] == 3
    assert log.exists()
